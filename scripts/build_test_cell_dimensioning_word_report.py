#!/usr/bin/env python3
"""Build the authoritative Word report for the integrated sealed EO cell concept."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "config" / "eo_vapor_test_cell_design.json"
MODEL_PATH = ROOT / "scripts" / "design_eo_vapor_test_cell.py"
OUTPUT_DIR = ROOT / "reports" / "test_cell_handoff_supplement_2026-09-03"
OUTPUT_FILE = (
    OUTPUT_DIR
    / "HeiQ_Integrated_EO_Vapour_Test_Cell_Handoff_2026-09-03_v1.docx"
)
FIGURE = OUTPUT_DIR / "figures" / "eo_vapour_integrated_two_chamber_dimensioned_views.png"


def load_model():
    spec = importlib.util.spec_from_file_location("eo_cell_model", MODEL_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load {MODEL_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    font_size_pt: float = 8.5,
    column_widths_cm: list[float] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    if column_widths_cm is not None:
        table.autofit = False
    for row_number, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_number == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_number, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if column_widths_cm is not None:
                cell.width = Cm(column_widths_cm[column_number])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size_pt)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Heading 1"]
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    model = load_model()
    volumes = model.build_integrated_volume_basis_frame(config)
    transient = model.build_integrated_transient_frame(config)
    register = model.build_integrated_engineering_register_frame(config)
    release = config["integrated_sealed_cell"]["fabrication_release"]
    basis = volumes.iloc[0]

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    document.core_properties.title = "Integrated Sealed EO Vapour Test Cell — Dimensioning Basis"
    document.core_properties.subject = "0.500 L-per-side engineering concept and transient inventory basis"
    document.core_properties.author = "HeiQ XMembrane Study Team"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(3)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Integrated Sealed EO Vapour Test Cell")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = document.add_paragraph("Dimensioning basis — current handoff release, 2026-09-03")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status = document.add_paragraph(
        "Status: authoritative engineering concept; fabrication release is blocked "
        "by the open evidence gates."
    )
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_table(
        document,
        ["NET-VOLUME TARGET / SIDE", "MEMBRANE REGION", "LOCAL GAPS"],
        [["500.0 +/- 5.0 mL", "50 x 50 mm (25 cm2)", "1 / 2 / 4 mm"]],
    )

    add_heading(document, "Selected physical system")
    document.add_paragraph(
        "Two equal sealed gas chambers, A and B, share the active membrane boundary. "
        "Each assembled chamber shall be calibrated to 500.0 +/- 5.0 mL net "
        "free-gas volume at "
        "32 °C, 101325 Pa and 50% initial RH. The membrane coupon is 80 x 80 mm, "
        "the active region is 50 x 50 mm, the clear chamber face is 100 x 100 mm, "
        "and the outer plate envelope is 110 x 110 mm. The volume is a governing "
        "design target; traceable assembled-cell calibration remains open."
    )
    if FIGURE.is_file():
        document.add_picture(str(FIGURE), width=Cm(16.2))
        caption = document.add_paragraph("Figure 1. Dimensioned exterior and section engineering concept.")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "Volume derived from target drift and switching time")
    document.add_paragraph(
        f"The declared target flux is 208 g m-2 h-1 across 25 cm2, giving "
        f"{basis['target_net_water_rate_mg_h']:.1f} mg/h ({basis['target_net_water_rate_mg_s']:.6f} mg/s). "
        f"A 5.0 s polarity segment transfers {basis['water_transferred_per_half_cycle_mg']:.6f} mg. "
        f"Using V_required = transferred mass / (saturated vapour density x 0.05), "
        f"the minimum net free-gas volume is {basis['required_net_free_gas_volume_ml_per_side']:.1f} mL per side. "
        f"The selected 500.0 mL target gives a ratio of {basis['selected_to_required_volume_ratio']:.3f} and "
        f"a predicted change of {basis['predicted_single_chamber_rh_change_percentage_points']:.3f} "
        "RH percentage points in either chamber at the first switching point."
    )

    add_heading(document, "Local, bulk, net and gross volume")
    add_table(
        document,
        ["Gap", "Local membrane region", "Bulk free gas", "Net-volume target"],
        [
            [
                f"{row.local_gap_mm:g} mm",
                f"{row.local_membrane_region_volume_ml_per_side:.1f} mL",
                f"{row.bulk_free_gas_volume_ml_per_side:.1f} mL",
                f"{row.calibrated_net_free_gas_volume_ml_per_side:.1f} mL",
            ]
            for row in volumes.itertuples()
        ],
    )
    volume_definition = document.add_paragraph(
        "The local membrane-region volume is included in the net-volume target. "
        "Gross empty-cavity volume is released from gross = calibrated net + installed "
        "component displacement after the internal hardware inventory is measured. "
        "The rendering does not define any dimension stored as null in the configuration."
    )
    volume_definition.paragraph_format.keep_together = True

    add_heading(document, "Ideal ABBA transient")
    endpoints = transient[transient["time_s"].isin([0.0, 5.0, 10.0, 15.0, 20.0])]
    add_table(
        document,
        ["Time", "RH A", "RH B", "Transfer A to B"],
        [
            [
                f"{row.time_s:.0f} s",
                f"{row.relative_humidity_A_percent:.3f}%",
                f"{row.relative_humidity_B_percent:.3f}%",
                f"{row.signed_water_transferred_A_to_B_mg:+.6f} mg",
            ]
            for row in endpoints.itertuples()
        ],
    )
    document.add_paragraph(
        "The analytic +1, -1, -1, +1 sequence conserves total water inventory, gives "
        "equal and opposite chamber changes, and returns both chambers to the initial state "
        f"at 20 s. The gas-only pressure diagnostic reaches "
        f"{transient['gas_only_differential_pressure_B_minus_A_pa'].abs().max():.1f} Pa "
        f"at an extremum. Pressure management remains open pending verification against "
        f"the {config['integrated_sealed_cell']['differential_pressure_target_pa']:.1f} Pa target."
    )

    add_heading(document, "Engineering register for fabrication release")
    document.add_paragraph(
        f"Fabrication release ready: {str(bool(release['ready'])).lower()}. Every "
        "required register item must be closed with cited evidence before this may "
        "be set to true. The gates explicitly cover walls and divider, ports and "
        "feedthroughs, fan, baffles, fixture mounting, fasteners, datums, tolerances, "
        "gap uniformity, pressure management, mixing response and net-volume calibration."
    )
    add_table(
        document,
        ["Item", "Status", "Closure evidence"],
        [
            [str(r.item).replace("_", " ").title(), str(r.status).title(), str(r.closure_evidence)]
            for r in register.itertuples()
        ],
        font_size_pt=7.4,
        column_widths_cm=[5.4, 1.8, 9.6],
    )

    add_heading(document, "Reproduction and verification")
    document.add_paragraph(
        "Run scripts/design_eo_vapor_test_cell.py to regenerate the volume basis, 0.1 s "
        "transient, engineering register, audit and source manifest. Run "
        "pytest -q tests/test_eo_vapor_test_cell_design.py for the formula, conservation, "
        "symmetry and packaged-output checks."
    )

    document.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
