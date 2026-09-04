#!/usr/bin/env python3
"""Build a short English design-basis document for the EO vapour test cell."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ModuleNotFoundError:
    DOCX_AVAILABLE = False
else:
    DOCX_AVAILABLE = True


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "config" / "eo_vapor_test_cell_net_transport_volume_study.json"
BASE_CONFIG_PATH = ROOT / "config" / "eo_vapor_test_cell_design.json"
BASE_MODEL_PATH = ROOT / "scripts" / "design_eo_vapor_test_cell.py"
GEOMETRY_SUMMARY_PATH = (
    ROOT
    / "reports"
    / "test_cell_20mm_channel_study_2026-09-04"
    / "decision_summary.json"
)
DEFAULT_OUTPUT_DIR = (
    ROOT / "reports" / "test_cell_selected_design_basis_2026-09-04"
)
DOCUMENT_STEM = "HeiQ_EO_Vapour_Test_Cell_Selected_Design_Basis_2026-09-04"
BLUEPRINT_STEM = "eo_vapour_test_cell_engineering_blueprint"

if DOCX_AVAILABLE:
    NAVY = RGBColor(20, 48, 75)
    TEAL = RGBColor(23, 107, 135)
    ORANGE = RGBColor(217, 119, 6)
    SLATE = RGBColor(71, 85, 105)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--skip-docx", action="store_true")
    parser.add_argument("--skip-pdf", action="store_true")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_base_model():
    spec = importlib.util.spec_from_file_location("eo_cell_basis_model", BASE_MODEL_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load {BASE_MODEL_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def calculate_basis() -> dict[str, float | list[int] | str]:
    study = read_json(CONFIG_PATH)
    base = read_json(BASE_CONFIG_PATH)
    geometry = read_json(GEOMETRY_SUMMARY_PATH)
    model = load_base_model()

    selected = study["recommended_working_point"]
    net_volume_ml = float(selected["calibrated_net_volume_ml_per_side"])
    segment_s = float(selected["segment_duration_s"])
    tau_s = float(study["transient"]["sensor_tau63_s"])
    active_area_m2 = float(model.membrane_area_m2(base))
    active_area_cm2 = active_area_m2 * 1.0e4
    flux_g_m2_h = float(
        base["integrated_sealed_cell"]["volume_sizing_basis"][
            "target_net_water_flux_g_m2_h"
        ]
    )
    rate_mg_s = float(model.integrated_target_water_rate_mg_s(base))
    transferred_mg = rate_mg_s * segment_s
    temperature_c = float(base["operating_state"]["temperature_c"])
    temperature_k = temperature_c + 273.15
    absolute_pressure_pa = float(base["operating_state"]["absolute_pressure_pa"])
    initial_rh_percent = float(
        base["operating_state"]["initial_relative_humidity_percent"]
    )
    saturation_density_kg_m3 = float(
        model.saturated_water_vapour_density_kg_m3(temperature_c)
    )
    rh_change_pp = (
        transferred_mg
        / 1.0e6
        / (saturation_density_kg_m3 * net_volume_ml / 1.0e6)
        * 100.0
    )
    water_molar_mass_kg_mol = 0.01801528
    gas_constant_j_mol_k = 8.31446261815324
    transferred_moles = transferred_mg / 1.0e6 / water_molar_mass_kg_mol
    rigid_dp_pa = (
        2.0
        * transferred_moles
        * gas_constant_j_mol_k
        * temperature_k
        / (net_volume_ml / 1.0e6)
    )
    target_dp_pa = float(base["integrated_sealed_cell"]["differential_pressure_target_pa"])
    half_target_dp_pa = target_dp_pa / 2.0
    net_volume_m3 = net_volume_ml / 1.0e6
    nrt = transferred_moles * gas_constant_j_mol_k * temperature_k
    high_volume_m3 = (
        absolute_pressure_pa * net_volume_m3 + nrt
    ) / (absolute_pressure_pa + half_target_dp_pa)
    low_volume_m3 = (
        absolute_pressure_pa * net_volume_m3 - nrt
    ) / (absolute_pressure_pa - half_target_dp_pa)
    diaphragm_travel_ml = max(
        abs(high_volume_m3 - net_volume_m3),
        abs(low_volume_m3 - net_volume_m3),
    ) * 1.0e6
    diaphragm_full_stroke_ml = (high_volume_m3 - low_volume_m3) * 1.0e6
    sensor_step_fraction = 1.0 - math.exp(-segment_s / tau_s)
    sensor_ramp_fraction = 1.0 - tau_s / segment_s * sensor_step_fraction

    recommended_geometry = geometry["recommended_geometry"]
    return {
        "study_id": study["study_id"],
        "active_area_m2": active_area_m2,
        "active_area_cm2": active_area_cm2,
        "active_width_mm": float(study["fixed_membrane"]["active_window_mm"][0]),
        "active_height_mm": float(study["fixed_membrane"]["active_window_mm"][1]),
        "coupon_width_mm": float(study["fixed_membrane"]["coupon_mm"][0]),
        "coupon_height_mm": float(study["fixed_membrane"]["coupon_mm"][1]),
        "local_gap_mm": float(study["fixed_membrane"]["local_membrane_gap_mm"]),
        "local_gap_volume_ml": active_area_m2 * float(study["fixed_membrane"]["local_membrane_gap_mm"]) * 1000.0,
        "bulk_depth_mm": float(study["bulk_geometry"]["bulk_depth_mm"]),
        "bulk_footprint_width_mm": float(recommended_geometry["bulk_cavity_footprint_mm"][0]),
        "bulk_footprint_height_mm": float(recommended_geometry["bulk_cavity_footprint_mm"][1]),
        "gross_volume_ml": float(recommended_geometry["gross_volume_ml_per_side"]),
        "component_bound_ml": float(recommended_geometry["component_displacement_upper_bound_ml_per_side"]),
        "pretrim_net_volume_ml": float(recommended_geometry["pretrim_net_volume_ml_per_side"]),
        "nominal_trim_ml": float(recommended_geometry["nominal_removable_trim_insert_ml_per_side"]),
        "maximum_trim_capacity_ml": float(recommended_geometry["minimum_available_trim_capacity_ml_per_side"]),
        "net_volume_ml": net_volume_ml,
        "volume_tolerance_ml": float(recommended_geometry["final_acceptance_tolerance_ml"]),
        "temperature_c": temperature_c,
        "absolute_pressure_pa": absolute_pressure_pa,
        "initial_rh_percent": initial_rh_percent,
        "flux_g_m2_h": flux_g_m2_h,
        "rate_mg_s": rate_mg_s,
        "rate_mg_h": rate_mg_s * 3600.0,
        "segment_s": segment_s,
        "polarity_sequence": list(selected["polarity_sequence"]),
        "abba_duration_s": segment_s * len(selected["polarity_sequence"]),
        "transferred_mg": transferred_mg,
        "saturation_density_kg_m3": saturation_density_kg_m3,
        "rh_change_pp": rh_change_pp,
        "rh_low_percent": initial_rh_percent - rh_change_pp,
        "rh_high_percent": initial_rh_percent + rh_change_pp,
        "differential_rh_pp": 2.0 * rh_change_pp,
        "rigid_dp_pa": rigid_dp_pa,
        "target_dp_pa": target_dp_pa,
        "diaphragm_travel_ml": diaphragm_travel_ml,
        "diaphragm_full_stroke_ml": diaphragm_full_stroke_ml,
        "sensor_tau_s": tau_s,
        "sensor_step_fraction": sensor_step_fraction,
        "sensor_ramp_fraction": sensor_ramp_fraction,
        "sensor_ramp_indicated_pp": rh_change_pp * sensor_ramp_fraction,
        "selection_reason": str(selected["selection_reason"]),
    }


def build_blueprint(png_path: Path, svg_path: Path) -> None:
    """Draw a number-free functional section and cartridge detail."""
    import matplotlib.pyplot as plt
    from matplotlib.patches import Arc, Circle, FancyArrowPatch, Rectangle

    ink = "#17304b"
    light = "#e9f2f8"
    fill = "#f8fbfd"
    accent = "#127a9b"
    orange = "#d97706"
    grey = "#65758b"

    figure, axis = plt.subplots(figsize=(13.5, 7.6))
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")
    axis.set_xlim(0, 16)
    axis.set_ylim(0, 9)
    axis.axis("off")

    def box(x, y, width, height, *, face=fill, linewidth=1.6, linestyle="-"):
        patch = Rectangle(
            (x, y), width, height, facecolor=face, edgecolor=ink,
            linewidth=linewidth, linestyle=linestyle, zorder=2,
        )
        axis.add_patch(patch)
        return patch

    def label(x, y, text, *, size=8.2, weight="normal", align="center", color=ink):
        axis.text(
            x, y, text, ha=align, va="center", fontsize=size,
            fontweight=weight, color=color, family="DejaVu Sans", zorder=8,
        )

    def leader(x_start, y_start, x_end, y_end, text, *, side="right"):
        axis.plot([x_start, x_end], [y_start, y_end], color=grey, linewidth=1.0, zorder=5)
        alignment = "left" if side == "right" else "right"
        offset = 0.08 if side == "right" else -0.08
        label(x_end + offset, y_end, text, size=7.4, align=alignment, color=grey)

    label(0.4, 8.62, "EO VAPOUR TEST CELL — FUNCTIONAL SECTION", size=13, weight="bold", align="left")
    axis.plot([0.4, 15.6], [8.33, 8.33], color=accent, linewidth=2.0)

    # Main symmetric chamber assembly.
    box(0.7, 3.15, 6.05, 3.85, face=light)
    box(9.25, 3.15, 6.05, 3.85, face=light)
    box(6.75, 2.9, 2.5, 4.35, face="#f1f5f9", linewidth=2.0)
    box(7.45, 3.15, 0.22, 3.85, face="#c7d7e3", linewidth=1.0)
    box(7.67, 3.15, 0.18, 3.85, face="#f3c27a", linewidth=1.0)
    box(7.85, 3.15, 0.30, 3.85, face="#ffffff", linewidth=1.3)
    box(8.15, 3.15, 0.18, 3.85, face="#f3c27a", linewidth=1.0)
    box(8.33, 3.15, 0.22, 3.85, face="#c7d7e3", linewidth=1.0)

    label(3.72, 6.65, "CONDITIONED CHAMBER — SIDE A", size=9.3, weight="bold")
    label(12.28, 6.65, "CONDITIONED CHAMBER — SIDE B", size=9.3, weight="bold")
    label(8.0, 7.55, "REMOVABLE MEMBRANE CARTRIDGE", size=8.6, weight="bold")

    # Fans and mixing loops.
    for fan_x in (1.55, 14.45):
        axis.add_patch(Circle((fan_x, 4.25), 0.48, facecolor="white", edgecolor=ink, linewidth=1.4, zorder=4))
        axis.add_patch(Circle((fan_x, 4.25), 0.08, facecolor=ink, edgecolor=ink, zorder=5))
        for angle in (0, 90, 180, 270):
            axis.add_patch(Arc((fan_x, 4.25), 0.72, 0.36, angle=angle, theta1=15, theta2=165, color=ink, linewidth=1.1, zorder=5))
    box(2.3, 3.65, 0.16, 2.15, face="white", linewidth=1.1)
    box(13.54, 3.65, 0.16, 2.15, face="white", linewidth=1.1)
    for x_start, x_end, y, bend in ((1.8, 6.45, 5.9, 0.35), (14.2, 9.55, 5.9, -0.35), (6.4, 1.85, 3.55, -0.25), (9.6, 14.15, 3.55, 0.25)):
        axis.add_patch(FancyArrowPatch(
            (x_start, y), (x_end, y), arrowstyle="-|>", mutation_scale=11,
            connectionstyle=f"arc3,rad={bend}", color=accent, linewidth=1.25, zorder=5,
        ))

    # Sensor probes and pressure taps.
    for probe_x in (4.15, 11.85):
        axis.plot([probe_x, probe_x], [7.0, 5.45], color=ink, linewidth=1.6, zorder=5)
        axis.add_patch(Circle((probe_x, 5.25), 0.20, facecolor="white", edgecolor=ink, linewidth=1.2, zorder=5))
    for tap_x in (5.35, 10.65):
        axis.plot([tap_x, tap_x], [7.0, 6.15], color=ink, linewidth=1.4, zorder=5)
        axis.add_patch(Circle((tap_x, 6.05), 0.09, facecolor=orange, edgecolor=ink, linewidth=0.8, zorder=6))

    # Flexible pressure compensators beneath each chamber.
    for cx in (4.7, 11.3):
        box(cx - 0.9, 2.35, 1.8, 0.48, face="white", linewidth=1.2)
        xs = [cx - 0.72, cx - 0.45, cx - 0.18, cx + 0.09, cx + 0.36, cx + 0.63]
        ys = [2.58, 2.43, 2.70, 2.43, 2.70, 2.58]
        axis.plot(xs, ys, color=accent, linewidth=1.2, zorder=5)
        axis.plot([cx, cx], [3.15, 2.83], color=ink, linewidth=1.0, zorder=4)

    # Measurement and electrical topology.
    box(6.1, 0.72, 3.8, 0.92, face="#eef4f7", linewidth=1.5)
    label(8.0, 1.18, "BIPOLAR SOURCE–MEASURE UNIT", size=8.6, weight="bold")
    axis.plot([7.75, 7.75], [1.64, 2.9], color=orange, linewidth=1.6, zorder=4)
    axis.plot([8.25, 8.25], [1.64, 2.9], color=ink, linewidth=1.6, zorder=4)
    axis.add_patch(FancyArrowPatch((7.38, 5.08), (8.62, 5.08), arrowstyle="<|-|>", mutation_scale=13, color=orange, linewidth=1.5, zorder=7))
    label(8.0, 4.72, "NET WATER TRANSPORT", size=7.2, weight="bold", color=orange)

    box(0.75, 0.72, 3.25, 0.92, face="#eef4f7", linewidth=1.3)
    label(2.38, 1.18, "DATA ACQUISITION", size=8.3, weight="bold")
    box(12.0, 0.72, 3.25, 0.92, face="#eef4f7", linewidth=1.3)
    label(13.62, 1.18, "DIFFERENTIAL PRESSURE", size=8.0, weight="bold")
    axis.plot([4.15, 2.4, 2.4], [7.0, 7.75, 1.64], color=grey, linewidth=1.0, linestyle="--", zorder=1)
    axis.plot([11.85, 2.8, 2.8], [7.0, 7.55, 1.64], color=grey, linewidth=1.0, linestyle="--", zorder=1)
    axis.plot([5.35, 13.2, 13.2], [7.0, 7.75, 1.64], color=grey, linewidth=1.0, linestyle="--", zorder=1)
    axis.plot([10.65, 14.0, 14.0], [7.0, 7.55, 1.64], color=grey, linewidth=1.0, linestyle="--", zorder=1)

    # Callouts.
    leader(4.15, 5.3, 3.2, 5.0, "temperature / RH probe", side="left")
    leader(5.35, 6.05, 5.85, 6.35, "pressure tap")
    leader(1.55, 4.25, 0.7, 2.95, "quiet mixing fan", side="right")
    leader(2.38, 4.6, 2.9, 4.25, "flow baffle")
    leader(4.7, 2.6, 3.35, 2.05, "flexible pressure compensator", side="left")
    leader(6.75, 6.7, 6.1, 7.72, "clamping frame", side="left")
    leader(7.56, 4.05, 6.85, 3.45, "local gas gap", side="left")
    leader(7.76, 5.65, 6.8, 5.85, "porous electrode", side="left")
    leader(8.0, 5.5, 9.0, 5.85, "EO membrane coupon")
    leader(8.45, 4.05, 9.15, 3.45, "local gas gap")
    leader(8.28, 6.2, 9.2, 6.4, "double seal / guard groove")

    label(0.7, 0.25, "Solid line: gas or electrical connection     Dashed line: instrumentation signal", size=7.4, align="left", color=grey)
    label(15.3, 0.25, "Functional layout — not a fabrication drawing", size=7.4, align="right", color=grey)

    png_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(png_path, dpi=220, bbox_inches="tight", facecolor="white")
    figure.savefig(svg_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def add_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        add_cell_shading(cell, "17304B")
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            set_cell_text(cell, text)
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Heading 1"]
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.color.rgb = NAVY


def add_equation(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10)
    run.font.color.rgb = TEAL


def build_markdown(values: dict[str, Any], output_path: Path) -> None:
    sequence = ", ".join(f"{value:+d}" for value in values["polarity_sequence"])
    text = f"""# EO vapour test cell — selected design basis

Status: engineering working basis for net-transient membrane measurements; not a fabrication-release drawing.

![Two-dimensional engineering blueprint of the test cell](figures/eo_vapour_test_cell_engineering_blueprint.png)

*Figure 1. Two-dimensional functional section showing the central membrane cartridge, symmetric conditioned chambers, sensing, internal mixing, pressure compensation, source-measure unit and acquisition connections. Values and dimensions are deliberately kept outside the figure. It is an engineering system blueprint, not released fabrication CAD.*

## Selected working values

| Quantity | Selected value | Reason |
| --- | ---: | --- |
| Active membrane window | {values['active_width_mm']:.0f} x {values['active_height_mm']:.0f} mm = {values['active_area_cm2']:.0f} cm² | Largest existing aperture, strong absolute water signal and lower edge-to-area ratio. |
| Membrane coupon | {values['coupon_width_mm']:.0f} x {values['coupon_height_mm']:.0f} mm | Leaves land for the double seal, guard groove and clamping. |
| Local gas gap | {values['local_gap_mm']:.0f} mm per side | Matches the local transport model and gives {values['local_gap_volume_ml']:.1f} mL active-gap volume per side. |
| Calibrated net gas volume | {values['net_volume_ml']:.0f} ± {values['volume_tolerance_ml']:.0f} mL per side | Retains the illustrated 0.50 L cell and slows the humidity transient enough for the present instrumentation. |
| Bulk plenum | {values['bulk_footprint_width_mm']:.0f} x {values['bulk_footprint_height_mm']:.0f} x {values['bulk_depth_mm']:.0f} mm | Produces {values['gross_volume_ml']:.1f} mL gross volume while keeping the active membrane fixed at 25 cm². |
| Polarity protocol | {values['segment_s']:.0f} s segments; {sequence} | Reports net transferred water per segment and uses reversal to suppress fixed and linear drift. |

## Calculation basis

At {values['temperature_c']:.0f} °C, {values['absolute_pressure_pa'] / 1000.0:.3f} kPa and {values['initial_rh_percent']:.0f}% initial RH, the conditional design flux is {values['flux_g_m2_h']:.0f} g m⁻² h⁻¹. This is a sizing target, not a measured membrane property.

**Integrated rate**

`ṁ = J A = {values['flux_g_m2_h']:.0f} x {values['active_area_m2']:.4f} = {values['rate_mg_h']:.0f} mg h⁻¹ = {values['rate_mg_s']:.6f} mg s⁻¹`

**Transferred mass in one segment**

`Δm = ṁ Δt = {values['rate_mg_s']:.6f} x {values['segment_s']:.0f} = {values['transferred_mg']:.6f} mg`

**Ideal well-mixed humidity response**

`ΔRH = Δm / (ρ_sat V) = {values['rh_change_pp']:.3f} percentage points per chamber`

The first positive segment therefore moves the ideal chamber endpoints from {values['initial_rh_percent']:.0f}%/{values['initial_rh_percent']:.0f}% to {values['rh_low_percent']:.3f}%/{values['rh_high_percent']:.3f}%. The chamber-to-chamber difference is {values['differential_rh_pp']:.3f} percentage points.

**Volume accounting**

`{values['gross_volume_ml']:.1f} mL gross − {values['component_bound_ml']:.1f} mL component bound = {values['pretrim_net_volume_ml']:.1f} mL pre-trim`

`{values['pretrim_net_volume_ml']:.1f} mL − {values['nominal_trim_ml']:.1f} mL nominal insert = {values['net_volume_ml']:.1f} mL calibrated net gas`

The removable trim system must cover {values['nominal_trim_ml']:.1f}–{values['maximum_trim_capacity_ml']:.1f} mL because actual solid displacement is not yet measured.

## Why this working point

The 500 mL inventory prevents the response from becoming excessively fast while retaining a clear net-transient signal. A {values['segment_s']:.0f} s segment gives approximately {values['rh_change_pp']:.1f} percentage points per chamber. The SHT45 step-response fraction over this duration is {values['sensor_step_fraction']:.1%}; however, a continuously changing RH ramp would only be indicated at about {values['sensor_ramp_fraction']:.1%} in a simple first-order sensor model. A faster calibrated reference remains preferable for the transient shape.

The rigid-cell gas diagnostic is {values['rigid_dp_pa']:.0f} Pa after one segment, far above the {values['target_dp_pa']:.0f} Pa differential-pressure target. Under the same ideal-gas screen, pressure management therefore needs approximately ±{values['diaphragm_travel_ml']:.2f} mL travel per side and {values['diaphragm_full_stroke_ml']:.2f} mL full stroke. This requirement must be verified with the real diaphragm, fan, seals and ports.

## Measurement interpretation

The reported quantity is integrated **net water movement through the membrane during each segment** under evolving RH conditions. It includes EO-driven transport together with opposing diffusion and other real membrane effects. The symmetric sequence lasts {values['abba_duration_s']:.0f} s and returns the ideal closed-cell inventory to its initial state after a complete block. It does not demonstrate sustained dehumidification.

The selected dimensions remain conditional on measured mixing, leakage, temperature bias, differential pressure, installed component displacement and calibrated net volume. The visualization is not fabrication evidence.
"""
    output_path.write_text(text, encoding="utf-8")


def build_docx(values: dict[str, Any], figure_path: Path, output_path: Path) -> None:
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "python-docx is required for DOCX output; install "
            "requirements-docs-presentations.txt or use --skip-docx --skip-pdf"
        )
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    document.core_properties.title = "EO Vapour Test Cell — Selected Design Basis"
    document.core_properties.subject = "Selected values and calculations for the net-transient two-chamber cell"
    document.core_properties.author = "HeiQ XMembrane Study Team"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Normal"].paragraph_format.space_after = Pt(3)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("EO Vapour Test Cell")
    title_run.bold = True
    title_run.font.size = Pt(21)
    title_run.font.color.rgb = NAVY
    subtitle = document.add_paragraph("Selected design basis for net-transient measurement")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = TEAL
    status = document.add_paragraph(
        "Engineering working basis · calculation-backed concept · not released for fabrication"
    )
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.runs[0].italic = True
    status.runs[0].font.color.rgb = SLATE

    add_table(
        document,
        ["ACTIVE MEMBRANE", "NET VOLUME / SIDE", "WORKING SEGMENT"],
        [[
            f"{values['active_width_mm']:.0f} x {values['active_height_mm']:.0f} mm ({values['active_area_cm2']:.0f} cm²)",
            f"{values['net_volume_ml']:.0f} ± {values['volume_tolerance_ml']:.0f} mL",
            f"{values['segment_s']:.0f} s",
        ]],
    )
    document.add_paragraph()
    document.add_picture(str(figure_path), width=Cm(17.0))
    caption = document.add_paragraph(
        "Figure 1. Two-dimensional functional blueprint of the complete test setup. Values and dimensions are kept in the calculation basis; this is not released fabrication CAD."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)
    caption.runs[0].font.color.rgb = SLATE

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "Selected values and reasons")
    add_table(
        document,
        ["Quantity", "Selected value", "Why it was selected"],
        [
            [
                "Active membrane",
                f"{values['active_width_mm']:.0f} x {values['active_height_mm']:.0f} mm = {values['active_area_cm2']:.0f} cm²",
                "Largest existing aperture; strong signal and lower edge-to-area ratio.",
            ],
            [
                "Membrane coupon",
                f"{values['coupon_width_mm']:.0f} x {values['coupon_height_mm']:.0f} mm",
                "Provides land for double seals, guard groove and clamping.",
            ],
            [
                "Local gas gap",
                f"{values['local_gap_mm']:.0f} mm per side",
                "Matches the local transport model; active-gap volume is only 5.0 mL per side.",
            ],
            [
                "Calibrated net gas volume",
                f"{values['net_volume_ml']:.0f} ± {values['volume_tolerance_ml']:.0f} mL per side",
                "Retains the illustrated 0.50 L cell and slows the transient for the present instrumentation.",
            ],
            [
                "Bulk plenum",
                f"{values['bulk_footprint_width_mm']:.0f} x {values['bulk_footprint_height_mm']:.0f} x {values['bulk_depth_mm']:.0f} mm",
                "Provides gross volume without changing the 25 cm² active membrane.",
            ],
            [
                "Polarity sequence",
                f"{values['segment_s']:.0f} s each; +I, −I, −I, +I",
                "Measures net transfer by polarity while suppressing fixed and linear drift.",
            ],
        ],
    )

    add_heading(document, "Calculation basis")
    document.add_paragraph(
        f"The calculation is conditional on {values['flux_g_m2_h']:.0f} g m⁻² h⁻¹ at "
        f"{values['temperature_c']:.0f} °C, {values['absolute_pressure_pa'] / 1000.0:.3f} kPa and "
        f"{values['initial_rh_percent']:.0f}% initial RH. The flux is a test-sizing target, not a measured material property."
    )
    add_equation(
        document,
        f"ṁ = J A = {values['flux_g_m2_h']:.0f} × {values['active_area_m2']:.4f} = "
        f"{values['rate_mg_h']:.0f} mg/h = {values['rate_mg_s']:.6f} mg/s",
    )
    add_equation(
        document,
        f"Δm = ṁ Δt = {values['rate_mg_s']:.6f} × {values['segment_s']:.0f} = "
        f"{values['transferred_mg']:.6f} mg",
    )
    add_equation(
        document,
        f"ΔRH = Δm / (ρsat Vnet) = {values['rh_change_pp']:.3f} percentage points per chamber",
    )
    document.add_paragraph(
        f"After the first positive segment, the ideal well-mixed endpoints are "
        f"{values['rh_low_percent']:.3f}% and {values['rh_high_percent']:.3f}% RH. "
        f"The differential RH signal is {values['differential_rh_pp']:.3f} percentage points."
    )

    add_heading(document, "Volume closure")
    add_equation(
        document,
        f"{values['gross_volume_ml']:.1f} mL gross − {values['component_bound_ml']:.1f} mL components "
        f"− {values['nominal_trim_ml']:.1f} mL trim = {values['net_volume_ml']:.1f} mL net",
    )
    document.add_paragraph(
        f"The gross {values['bulk_footprint_width_mm']:.0f} x {values['bulk_footprint_height_mm']:.0f} x "
        f"{values['bulk_depth_mm']:.0f} mm plenum is {values['gross_volume_ml']:.1f} mL. "
        f"After the conservative {values['component_bound_ml']:.1f} mL component bound, "
        f"{values['pretrim_net_volume_ml']:.1f} mL remains. Removable inserts must cover "
        f"{values['nominal_trim_ml']:.1f}–{values['maximum_trim_capacity_ml']:.1f} mL until actual "
        "solid displacement is measured."
    )

    add_heading(document, "Instrument and pressure implications")
    document.add_paragraph(
        f"The SHT45 step-response fraction over {values['segment_s']:.0f} s is "
        f"{values['sensor_step_fraction']:.1%}. This does not mean exact ramp tracking: a simple "
        f"first-order model indicates only {values['sensor_ramp_fraction']:.1%} of a linear endpoint "
        f"change. Use a faster calibrated reference if the transient shape is required. The rigid-cell "
        f"pressure diagnostic is {values['rigid_dp_pa']:.0f} Pa, compared with the {values['target_dp_pa']:.0f} Pa "
        f"target. The ideal pressure-management screen therefore requires about ±{values['diaphragm_travel_ml']:.2f} mL "
        f"travel per side and {values['diaphragm_full_stroke_ml']:.2f} mL full stroke."
    )

    add_heading(document, "Interpretation boundary")
    paragraph = document.add_paragraph()
    paragraph.add_run("Report: ").bold = True
    paragraph.add_run(
        "integrated net water moved through the membrane during each segment under evolving RH conditions. "
    )
    paragraph.add_run("Do not claim: ").bold = True
    paragraph.add_run(
        "constant-boundary intrinsic EO flux or sustained dehumidification from the closed symmetric sequence alone. "
        f"One four-segment block lasts {values['abba_duration_s']:.0f} s and returns the ideal inventory to its initial state."
    )
    note = document.add_paragraph(
        "Fabrication remains conditional on measured mixing, leakage, fan thermal bias, differential pressure, "
        "component displacement, seal performance and assembled net volume."
    )
    note.runs[0].bold = True
    note.runs[0].font.color.rgb = ORANGE

    document.save(output_path)


def convert_pdf(docx_path: Path, output_dir: Path) -> Path:
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if office is None:
        raise RuntimeError("LibreOffice is required to build the PDF")
    with tempfile.TemporaryDirectory(prefix="eo-cell-lo-") as profile_dir:
        profile_uri = Path(profile_dir).resolve().as_uri()
        subprocess.run(
            [
                office,
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise RuntimeError("LibreOffice did not produce the expected PDF")
    return pdf_path


def write_manifest(output_dir: Path, generated: list[Path]) -> None:
    sources = [
        CONFIG_PATH,
        BASE_CONFIG_PATH,
        BASE_MODEL_PATH,
        GEOMETRY_SUMMARY_PATH,
        Path(__file__).resolve(),
    ]
    manifest = {
        "schema_version": 1,
        "document_id": "eo-vapour-test-cell-selected-design-basis-2026-09-04",
        "sources": [
            {
                "path": path.relative_to(ROOT).as_posix(),
                "sha256": sha256(path),
            }
            for path in sources
        ],
        "generated_outputs": [
            {
                "path": path.relative_to(output_dir).as_posix(),
                "sha256": sha256(path),
                "size_bytes": path.stat().st_size,
            }
            for path in generated
        ],
    }
    (output_dir / "source_manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def main() -> None:
    args = parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    values = calculate_basis()
    calculation_path = output_dir / "calculation_basis.json"
    markdown_path = output_dir / f"{DOCUMENT_STEM}.md"
    docx_path = output_dir / f"{DOCUMENT_STEM}.docx"
    figure_dir = output_dir / "figures"
    blueprint_png_path = figure_dir / f"{BLUEPRINT_STEM}.png"
    blueprint_svg_path = figure_dir / f"{BLUEPRINT_STEM}.svg"
    calculation_path.write_text(
        json.dumps(values, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    build_blueprint(blueprint_png_path, blueprint_svg_path)
    build_markdown(values, markdown_path)
    generated = [
        calculation_path,
        markdown_path,
        blueprint_png_path,
        blueprint_svg_path,
    ]
    if not args.skip_docx:
        build_docx(values, blueprint_png_path, docx_path)
        generated.append(docx_path)
    elif not args.skip_pdf:
        raise ValueError("PDF output requires DOCX output; add --skip-pdf")
    pdf_path = None
    if not args.skip_pdf:
        pdf_path = convert_pdf(docx_path, output_dir)
        generated.append(pdf_path)
    write_manifest(output_dir, generated)
    print(
        json.dumps(
            {
                "status": "PASS",
                "docx": str(docx_path),
                "pdf": str(pdf_path) if pdf_path else None,
                "pages_target": "short two-page document",
            },
            indent=2,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
